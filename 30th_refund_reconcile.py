import openpyxl, re, unicodedata
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC='/root/.claude/uploads/bdb5aef9-44bc-525f-9a08-c12e52c21899/40fc97e9-30th_Refund_List.xlsx'
OUT='/home/user/claude-aidesign/30th_Refund_List_RECONCILED.xlsx'

DP={'ETB':1185,'Sylveon Ex Box':540,'Ex Tin 5107':600,'Ex Tin 4107':540,'Binder':870,
    '2-Pack Blister':225,'Knock Out':225,'Poster':480,'Tech Sticker':480}
SRP={'ETB':3950,'Sylveon Ex Box':1800,'Ex Tin 5107':2000,'Ex Tin 4107':1800,'Binder':2900,
     '2-Pack Blister':750,'Knock Out':750,'Poster':1600,'Tech Sticker':1600}

WIN={'Binder':['Maribel Corpuz','Gabriel Galang'],
 '2-Pack Blister':['Louis Miguel Sta. Rita'],
 'Knock Out':['Adrian Paul Baldonado'],
 'Ex Tin 5107':['Jomariz Hesita','Daisy Alvior','Kerstie Barillo'],
 'Tech Sticker':['Arianne Carranceja','Kim Saballegue','Brandon Coyiuto','Hannah Mariano'],
 'Sylveon Ex Box':['Albie Peralta','Arianne Carranceja','Leoflor Jimenez','Al Francis Tesorero'],
 'Poster':['Albie Peralta','Jared Chua','Michael Vivo','Joric Cabial','Mark Laurence Recio','Kim Saballegue','Hannah Mariano'],
 'ETB':['Renz Ryan Garcia','Derick Valdez','Rein Kevin Legarda','Albie Peralta','John David Laranang','Kris Becker','Francis Chen']}

# Solo (non-raffle) orders filled directly pro-rata at 1.16%, allocation = round(qty x 1.16%).
# Reconciles to Russel's control table: ETB solo 29, 2-Pack Blister solo 1, total solo 30.
SOLO_RATE=0.0116
SOLO_ORDERS=[('John Paulo Manuel','ETB',1000),
             ('Anton Canoy','ETB',850),
             ('Vlade Mark Navarro','ETB',500),
             ('Alexis Tiutan','ETB',80),
             ('Gillian Mesoza','ETB',17),
             ('Joric Cabial','ETB',17),
             ('Lawrence (Anti-Hero)','2-Pack Blister',120)]
SOLO_ALLOC={}
for _n,_p,_q in SOLO_ORDERS:
    _a=round(_q*SOLO_RATE)
    if _a: SOLO_ALLOC.setdefault(_n,{})[_p]=_a

# Russel's control table (image, 2026-09-18): product -> (to_allocate, solo, raffle)
CONTROL={'ETB':(36,29,7),'2-Pack Blister':(2,1,1),'Knock Out':(1,0,1),'Sylveon Ex Box':(4,0,4),
         'Poster':(7,0,7),'Binder':(2,0,2),'Ex Tin 5107':(3,0,3),'Tech Sticker':(4,0,4),'Ex Tin 4107':(0,0,0)}

def norm(s):
    s=unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode()
    return set(re.sub(r'[^a-z ]','',s.lower()).split())
def money(v):
    if v is None or v=='': return None
    if isinstance(v,(int,float)): return float(v)
    try: return float(str(v).replace('₱','').replace(',','').strip())
    except: return None

wb=openpyxl.load_workbook(SRC,data_only=True); ws=wb['30th Refunds']
recs=[]
for r in range(4,174):
    name=ws.cell(r,2).value
    if not name: continue
    items={}
    for part in str(ws.cell(r,4).value or '').split(','):
        m=re.match(r'\s*(.+?)\s*x\s*(\d+)\s*$',part)
        if m: items[m.group(1)]=items.get(m.group(1),0)+int(m.group(2))
    recs.append(dict(src=r,name=str(name).strip(),oid=ws.cell(r,3).value,
        prods=ws.cell(r,4).value,items=items,qty=ws.cell(r,5).value,
        F=money(ws.cell(r,6).value),G=money(ws.cell(r,7).value),
        jacob=ws.cell(r,10).value,bank=ws.cell(r,11).value,acct=ws.cell(r,12).value,
        det=ws.cell(r,13).value,qr=ws.cell(r,14).value,done=ws.cell(r,15).value,
        alloc={},bulk=False))

# allocate: winner -> the one order row containing that product
unmatched=[]
for prod,names in WIN.items():
    for w in names:
        wt=norm(w)
        cands=[d for d in recs if wt<=norm(d['name']) or norm(d['name'])<=wt]
        hit=[d for d in cands if prod in d['items']]
        if len(hit)==1: hit[0]['alloc'][prod]=hit[0]['alloc'].get(prod,0)+1
        else: unmatched.append((prod,w,[d['src'] for d in hit]))
assert not unmatched, unmatched

for wname,alloc in SOLO_ALLOC.items():
    wt=norm(wname)
    for prod,qn in alloc.items():
        srcq=[q for n,p2,q in SOLO_ORDERS if n==wname and p2==prod][0]
        hit=[d for d in recs if (wt<=norm(d['name']) or norm(d['name'])<=wt)
             and d['items'].get(prod)==srcq]
        assert len(hit)==1, (wname,prod,[d['src'] for d in hit])
        hit[0]['alloc'][prod]=hit[0]['alloc'].get(prod,0)+qn
        hit[0]['bulk']=True

for d in recs:
    d['allocQty']=sum(d['alloc'].values())
    d['allocVal']=sum(SRP[k]*v for k,v in d['alloc'].items())
    d['paid']=d['G'] if d['G'] is not None else d['F']
    d['est']= d['G'] is None
    d['sfee']= (d['G']-d['F']) if d['G'] is not None else None
    d['net']=d['allocVal']-d['paid']            # +ve = customer owes, -ve = refund due
    d['refund']=-d['net'] if d['net']<0 else 0.0
    d['due']=d['net'] if d['net']>0 else 0.0

# ---------- write ----------
out=openpyxl.Workbook(); s=out.active; s.title='Refund Computation'
NAVY='1F3864'; HDR=PatternFill('solid',fgColor=NAVY)
WINF=PatternFill('solid',fgColor='FFF2CC'); MISS=PatternFill('solid',fgColor='FFC7CE')
BULK=PatternFill('solid',fgColor='FCE4D6'); DUEF=PatternFill('solid',fgColor='C6E0B4')
thin=Side(style='thin',color='BFBFBF'); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
P='₱#,##0.00'

s['A1']="30th Anniversary Pre-Order — Refund Computation AFTER Allocation"
s['A1'].font=Font(size=14,bold=True,color=NAVY)
s['A2']="Net = (SRP × Allocated Qty) − (Total 30% Downpayment with SFee).  Negative → REFUND to customer.  Positive → BALANCE DUE from customer."
s['A2'].font=Font(size=9,italic=True,color='808080')

H=['#','Name','Order ID','Products Ordered','Total Qty','30% DP (₱)','SFee (₱)',
   'Total 30% DP w/ SFee','Allocation Won','Alloc Qty','Allocation Value (SRP)',
   'NET (X − DP w/SFee)','REFUND DUE (₱)','BALANCE DUE (₱)','Status',
   'Cross-checked Jacob?','Bank Name','Account Owner','Bank / GCash No.','QR','Refunded? - Russ']
for i,h in enumerate(H,1):
    c=s.cell(4,i,h); c.font=Font(bold=True,color='FFFFFF',size=10); c.fill=HDR
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BD
s.freeze_panes='C5'; s.row_dimensions[4].height=34

row=5
for i,d in enumerate(recs,1):
    allocs=', '.join(f'{k} x{v}' for k,v in sorted(d['alloc'].items())) or '—'
    if d['est']: status='⚠ SFee NOT LOGGED — provisional (DP only)'
    elif d['allocQty']: status='WINNER — allocation applied'
    else: status='Full refund'
    vals=[i,d['name'],d['oid'],d['prods'],d['qty'],d['F'],d['sfee'],d['paid'],allocs,
          d['allocQty'] or None,d['allocVal'] or None,d['net'],d['refund'] or None,
          d['due'] or None,status,d['jacob'],d['bank'],d['acct'],d['det'],d['qr'],d['done']]
    for j,v in enumerate(vals,1):
        c=s.cell(row,j,v); c.border=BD; c.font=Font(size=10)
        if j in (6,7,8,11,12,13,14): c.number_format=P
        if j in (5,10): c.alignment=Alignment(horizontal='center')
        if j in (4,9,15): c.alignment=Alignment(wrap_text=True,vertical='top')
    if d['est']:
        for j in (7,8,12,13,14,15): s.cell(row,j).fill=MISS
    if d['allocQty']:
        for j in range(1,16): 
            if not (d['est'] and j in (7,8,12,13,14,15)): s.cell(row,j).fill=WINF
    if d['due']>0: s.cell(row,14).fill=DUEF; s.cell(row,14).font=Font(size=10,bold=True)
    if (d['qty'] or 0)>=30: s.cell(row,5).fill=BULK; s.cell(row,5).font=Font(size=10,bold=True,color='C00000')
    row+=1

t=row
s.cell(t,2,'TOTAL').font=Font(bold=True,size=11)
for col in (5,6,8,10,11,12,13,14):
    L=get_column_letter(col)
    c=s.cell(t,col,f'=SUM({L}5:{L}{row-1})')
    c.font=Font(bold=True,size=11); c.border=BD
    if col!=5 and col!=10: c.number_format=P
for j in range(1,22): s.cell(t,j).fill=PatternFill('solid',fgColor='D9E2F3'); s.cell(t,j).border=BD

for col,w in zip(range(1,22),[5,30,14,46,8,13,10,15,26,7,14,15,15,14,34,11,14,24,22,8,12]):
    s.column_dimensions[get_column_letter(col)].width=w
s.auto_filter.ref=f'A4:U{row-1}'

# ---- sheet 2: winners ----
s2=out.create_sheet('Winners — Allocation')
s2['A1']='Raffle Winners — Allocation Settlement'; s2['A1'].font=Font(size=14,bold=True,color=NAVY)
H2=['Product Won','Winner / Bulk Buyer','Order ID','SRP (₱)','Alloc Qty','Allocation Value',
    'Total 30% DP w/ SFee (whole order)','NET','Outcome','Status']
for i,h in enumerate(H2,1):
    c=s2.cell(3,i,h); c.font=Font(bold=True,color='FFFFFF',size=10); c.fill=HDR
    c.alignment=Alignment(horizontal='center',wrap_text=True); c.border=BD
s2.row_dimensions[3].height=32
r2=4
ORDER=[('ETB',[n for n in SOLO_ALLOC if 'ETB' in SOLO_ALLOC[n]]+WIN['ETB'])]+[(p,WIN[p]) for p in
   ['Binder','Ex Tin 5107','Sylveon Ex Box','Poster','Tech Sticker','2-Pack Blister','Knock Out']]
ORDER=[(p,([n for n in SOLO_ALLOC if p in SOLO_ALLOC[n]] if p=='2-Pack Blister' else [])+ns)
       if p=='2-Pack Blister' else (p,ns) for p,ns in ORDER]
for prod,names in ORDER:
    for w in names:
        wt=norm(w)
        d=[x for x in recs if (wt<=norm(x['name']) or norm(x['name'])<=wt) and prod in x['items']][0]
        out_txt=('REFUND ₱%s'%format(d['refund'],',.2f')) if d['refund']>0 else ('COLLECT ₱%s'%format(d['due'],',.2f'))
        st='⚠ SFee not logged — provisional' if d['est'] else 'OK'
        st=('SOLO pro-rata 1.16% \u2014 '+st) if d['bulk'] else st
        for j,v in enumerate([prod,d['name'],d['oid'],SRP[prod],d['alloc'][prod],
                              SRP[prod]*d['alloc'][prod],d['paid'],d['net'],out_txt,st],1):
            c=s2.cell(r2,j,v); c.border=BD; c.font=Font(size=10)
            if j in (4,6,7,8): c.number_format=P
            if j==5: c.alignment=Alignment(horizontal='center')
        if d['est']:
            for j in (7,8,9,10): s2.cell(r2,j).fill=MISS
        if d['due']>0: s2.cell(r2,9).fill=DUEF
        r2+=1
for col,w in zip(range(1,11),[17,26,14,12,10,16,22,14,22,30]):
    s2.column_dimensions[get_column_letter(col)].width=w

# ---- sheet 3: exceptions ----
s3=out.create_sheet('Exceptions')
s3['A1']='Items requiring action before payout'; s3['A1'].font=Font(size=14,bold=True,color='C00000')
H3=['Severity','Issue','Sheet Row','Name','Order ID','Detail']
for i,h in enumerate(H3,1):
    c=s3.cell(3,i,h); c.font=Font(bold=True,color='FFFFFF',size=10); c.fill=HDR; c.border=BD
r3=4
def ex(sev,issue,d,detail):
    global r3
    for j,v in enumerate([sev,issue,d['src'] if d else '',d['name'] if d else '',d['oid'] if d else '',detail],1):
        c=s3.cell(r3,j,v); c.border=BD; c.font=Font(size=10); c.alignment=Alignment(wrap_text=True,vertical='top')
    if sev=='BLOCKER': s3.cell(r3,1).fill=MISS
    elif sev=='VERIFY': s3.cell(r3,1).fill=BULK
    r3+=1
for d in recs:
    if d['est']: ex('BLOCKER','SFee not logged (Total 30% DP w/ SFee blank)',d,
        'Provisional figure uses 30%% DP only (₱%s). Add shipping fee before payout.'%format(d['F'],',.2f'))
for d in recs:
    if (d['qty'] or 0)>=30 and not d['bulk']: ex('VERIFY','Large order — raffle fill only',d,
        'Qty %s, DP ₱%s. Received no pro-rata allocation; confirm none is owed.'%(int(d['qty']),format(d['F'],',.2f')))
for d in recs:
    if d['bulk']: ex('BLOCKER','Very large refund — verify funds before release',d,
        'Pro-rata %s on %s ordered. Refund ₱%s. Confirm the ₱%s downpayment actually cleared and verify payee identity before releasing.'
        %(', '.join(f'{k} x{v}' for k,v in d['alloc'].items()),int(d['qty']),format(d['refund'],',.2f'),format(d['F'],',.2f')))
for d in recs:
    if d['G'] is not None and abs(d['G']-d['F'])<0.01: ex('VERIFY','SFee = ₱0',d,'Confirm free shipping / pickup.')
for d in recs:
    if str(d['oid']) in ('(none)','None','Meta','Meet'): ex('VERIFY','No Shopify order ID',d,'Order ID recorded as "%s" — manual/off-platform order.'%d['oid'])
for d in recs:
    if d['refund']>0 and not d['bank'] and not d['det']: ex('BLOCKER','No payout details on file',d,
        'Refund ₱%s owed but no bank/GCash recorded.'%format(d['refund'],',.2f'))
for col,w in zip(range(1,7),[12,42,10,28,14,72]): s3.column_dimensions[get_column_letter(col)].width=w
s3.auto_filter.ref=f'A3:F{r3-1}'

# ---- sheet 4: control reconciliation ----
s4=out.create_sheet('Control Check',0)
s4['A1']='Reconciliation vs Russel\u2019s Allocation Control Table'
s4['A1'].font=Font(size=14,bold=True,color=NAVY)
s4['A2']='Every figure below is recomputed from the refund sheet and compared against the control table.'
s4['A2'].font=Font(size=9,italic=True,color='808080')
H4=['Product','To Allocate (ctrl)','Solo (ctrl)','Raffle (ctrl)','Solo (computed)',
    'Raffle (computed)','Total (computed)','Match?','SRP (\u20b1)','Allocation Value (\u20b1)']
for i,h in enumerate(H4,1):
    c=s4.cell(4,i,h); c.font=Font(bold=True,color='FFFFFF',size=10); c.fill=HDR
    c.alignment=Alignment(horizontal='center',wrap_text=True); c.border=BD
s4.row_dimensions[4].height=34
solo_c={}; raf_c={}
for d in recs:
    for k,v in d['alloc'].items():
        (solo_c if d['bulk'] else raf_c).setdefault(k,0)
        if d['bulk']: solo_c[k]=solo_c.get(k,0)+v
        else: raf_c[k]=raf_c.get(k,0)+v
r4=5; allok=True
for prod,(ta,so,ra) in CONTROL.items():
    sc=solo_c.get(prod,0); rc=raf_c.get(prod,0); tc=sc+rc
    ok = (sc==so and rc==ra and tc==ta); allok = allok and ok
    for j,v in enumerate([prod,ta,so,ra,sc,rc,tc,'\u2713 MATCH' if ok else '\u2717 DIFF',
                          SRP[prod],SRP[prod]*tc],1):
        c=s4.cell(r4,j,v); c.border=BD; c.font=Font(size=10)
        if j in (2,3,4,5,6,7): c.alignment=Alignment(horizontal='center')
        if j in (9,10): c.number_format=P
    s4.cell(r4,8).fill=PatternFill('solid',fgColor='C6E0B4' if ok else 'FFC7CE')
    s4.cell(r4,8).font=Font(size=10,bold=True)
    r4+=1
for j,v in enumerate(['TOTAL',sum(v[0] for v in CONTROL.values()),sum(v[1] for v in CONTROL.values()),
                      sum(v[2] for v in CONTROL.values()),sum(solo_c.values()),sum(raf_c.values()),
                      sum(solo_c.values())+sum(raf_c.values()),
                      '\u2713 ALL MATCH' if allok else '\u2717 CHECK',
                      None,sum(SRP[k]*v for k,v in solo_c.items())+sum(SRP[k]*v for k,v in raf_c.items())],1):
    c=s4.cell(r4,j,v); c.border=BD; c.font=Font(size=11,bold=True)
    c.fill=PatternFill('solid',fgColor='D9E2F3')
    if j in (2,3,4,5,6,7,8): c.alignment=Alignment(horizontal='center')
    if j==10: c.number_format=P

s4.cell(r4+2,1,'Solo allocation rule: round(order qty \u00d7 1.16%). Reproduces both control totals (ETB solo 29, 2-Pack Blister solo 1) exactly.')
s4.cell(r4+2,1).font=Font(size=10,bold=True)
r5=r4+3
s4.cell(r5,1,'Solo order'); s4.cell(r5,2,'Product'); s4.cell(r5,3,'Qty ordered')
s4.cell(r5,4,'\u00d7 1.16%'); s4.cell(r5,5,'Allocated')
for j in range(1,6):
    c=s4.cell(r5,j); c.font=Font(bold=True,color='FFFFFF',size=10); c.fill=HDR; c.border=BD
r5+=1
for n,prod,q in SOLO_ORDERS:
    for j,v in enumerate([n,prod,q,round(q*SOLO_RATE,3),round(q*SOLO_RATE)],1):
        c=s4.cell(r5,j,v); c.border=BD; c.font=Font(size=10)
        if j in (3,4,5): c.alignment=Alignment(horizontal='center')
    if round(q*SOLO_RATE)==0: 
        for j in range(1,6): s4.cell(r5,j).font=Font(size=10,color='808080')
    r5+=1
s4.cell(r5,1,'Alexis Tiutan\u2019s 1 ETB unit is DERIVED, not supplied \u2014 confirm against your Solo Orders tab.')
s4.cell(r5,1).font=Font(size=10,italic=True,color='C00000')
for col,w in zip(range(1,11),[24,17,12,13,15,17,17,14,13,20]):
    s4.column_dimensions[get_column_letter(col)].width=w

out.save(OUT)

# ---------- console summary ----------
tot_paid=sum(d['paid'] for d in recs); tot_alloc=sum(d['allocVal'] for d in recs)
tot_ref=sum(d['refund'] for d in recs); tot_due=sum(d['due'] for d in recs)
miss=[d for d in recs if d['est']]
ref_known=sum(d['refund'] for d in recs if not d['est'])
ref_prov=sum(d['refund'] for d in recs if d['est'])
print('CONTROL TABLE CHECK:', 'ALL MATCH' if allok else 'MISMATCH')
print(f'orders                : {len(recs)}')
print(f'allocation units      : {sum(d["allocQty"] for d in recs)} across {len([d for d in recs if d["allocQty"]])} orders')
print(f'allocation value (SRP): {tot_alloc:>14,.2f}')
print(f'total DP collected    : {tot_paid:>14,.2f}  (incl. {len(miss)} rows w/o SFee)')
print(f'TOTAL REFUND          : {tot_ref:>14,.2f}   (confirmed {ref_known:,.2f} + provisional {ref_prov:,.2f})')
print(f'TOTAL BALANCE DUE     : {tot_due:>14,.2f}')
print(f'rows missing SFee     : {len(miss)}')
print(f'\nWinners still owing money:')
for d in recs:
    if d['due']>0: print(f'  {d["name"]:<26} {d["oid"]:<14} owes {d["due"]:>10,.2f}  ({", ".join(f"{k} x{v}" for k,v in d["alloc"].items())})')
print(f'\nWinners still getting a refund:')
for d in recs:
    if d['allocQty'] and d['refund']>0: print(f'  {d["name"]:<26} {d["oid"]:<14} refund {d["refund"]:>10,.2f}')
print('\nSaved:',OUT)
